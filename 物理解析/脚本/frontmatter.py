#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
frontmatter.py — 解析册的封面 / 说明页 / 版本记录页（生物、物理共用）

结构（5 节）：
  节1 封面      1 栏 · 垂直居中 · 无页码
  节2 说明页    1 栏 · 无页码
  节3 目录      2 栏 · 无页码
  节4 正文      2 栏 · 页码从 1 重新开始 · 奇数页右下 / 偶数页左下
  节5 版本记录  1 栏 · 无页码

关键技术点（均已实测验证）：
  ① `w:evenAndOddHeaders` 必须写在 **settings.xml**，不是 sectPr 的属性。
     写错位置时偶数页页脚永不生效（所有页码都跑右边）。
  ② 页码重启用 `<w:pgNumType w:start="1"/>`，且 sectPr 子元素必须按 schema 顺序插入。
  ③ WPS 的 `Information(1)`（adjusted）对分节后的页不可靠，一律返回 1；
     `Information(3)` 返回物理页。故显示页 = 物理页 − 前置页数偏移，
     偏移量由「正文首个标题的物理页 − 1」动态求得。
"""

from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm
from lxml import etree

FONT = '宋体'

# sectPr 子元素的合法顺序（OOXML CT_SectPr sequence）
_ORDER = ['headerReference', 'footerReference', 'footnotePr', 'endnotePr', 'type',
          'pgSz', 'pgMar', 'paperSrc', 'pgBorders', 'lnNumType', 'pgNumType', 'cols',
          'formProt', 'vAlign', 'noEndnote', 'titlePg', 'textDirection', 'bidi',
          'rtlGutter', 'docGrid', 'printerSettings']

# ── 文案（改文案只改这里）──
BRAND_MAIN = '全 解 全 析 整 理'
BRAND_DATE = '2026 年 9 月'
COLOPHON_TITLE = '本 册 说 明'
COLOPHON_LINES = [
    '整理排版：高三25班热心同学',
    '整理日期：2026 年 9 月',
    '内容来源：《步步高 大一轮》教师用书（{subject}·广东版）',
    '',
    '免责声明：内容版权归原出版方所有。本册仅为个人学习而做的版式整理，',
    '不作商业用途，也不对内容主张任何权利。',
    '不能保证 100% 正确，请辩证使用。',
    '如发现错误，请及时联系我更正。',
    '',
    '问题反馈（GitHub 公开仓库）：',
    'https://github.com/JunPulsar/Review-Analysis',
]
VERSION_TITLE = '版 本 记 录'
GITHUB_URL = 'https://github.com/JunPulsar/Review-Analysis'


# ────────────────────────── 底层工具 ──────────────────────────

def sect_child(sectPr, tag, **attrs):
    """按 schema 顺序插入/更新 sectPr 的子元素"""
    el = sectPr.find(qn('w:' + tag))
    if el is None:
        el = OxmlElement('w:' + tag)
        idx = _ORDER.index(tag)
        pos = None
        for i, ch in enumerate(sectPr):
            ct = ch.tag.split('}')[-1]
            if ct in _ORDER and _ORDER.index(ct) > idx:
                pos = i
                break
        if pos is None:
            sectPr.append(el)
        else:
            sectPr.insert(pos, el)
    for k, v in attrs.items():
        el.set(qn('w:' + k), v)
    return el


def set_font(run, name=FONT, size=None, bold=False, color=None):
    run.font.name = name
    run.font.bold = bold
    if size:
        run.font.size = size
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = etree.SubElement(rPr, qn('w:rFonts'))
    rFonts.set(qn('w:eastAsia'), name)
    return run


def enable_even_odd_headers(doc):
    """偶数页页眉页脚开关 —— 必须在 settings.xml（写成 sectPr 属性是无效的）"""
    settings = doc.settings.element
    if settings.find(qn('w:evenAndOddHeaders')) is None:
        settings.append(OxmlElement('w:evenAndOddHeaders'))


def _page_field(paragraph, align, size=Pt(12), prefix=''):
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if prefix:
        set_font(paragraph.add_run(prefix), size=size)
    r = paragraph.add_run()
    set_font(r, size=size)
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._element.append(b)
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = ' PAGE '
    r._element.append(it)
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._element.append(e)


def _blank_footer(section):
    section.footer.is_linked_to_previous = False
    for p in section.footer.paragraphs:
        p.clear()
    section.even_page_footer.is_linked_to_previous = False
    for p in section.even_page_footer.paragraphs:
        p.clear()


def _numbered_footer(section, footer_text=''):
    """奇数页右下 / 偶数页左下（外侧），可选左侧小字"""
    section.footer.is_linked_to_previous = False
    for p in section.footer.paragraphs:
        p.clear()
    _page_field(section.footer.paragraphs[0], WD_ALIGN_PARAGRAPH.RIGHT,
                prefix=(footer_text + '\t') if footer_text else '')
    section.even_page_footer.is_linked_to_previous = False
    for p in section.even_page_footer.paragraphs:
        p.clear()
    _page_field(section.even_page_footer.paragraphs[0], WD_ALIGN_PARAGRAPH.LEFT,
                prefix=('\t' + footer_text) if footer_text else '')


def _para(doc, text, size, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER,
          before=Pt(0), after=Pt(0), line=1.0, color=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = before
    p.paragraph_format.space_after = after
    p.paragraph_format.line_spacing = line
    if text:
        set_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def new_page_section(doc, cols=1, space_cm=0.6, vcenter=False, numbered=False,
                     footer_text='', restart=None):
    """新建一节并设置分栏 / 垂直居中 / 页码

    注意：python-docx 的 add_section 会克隆上一节的 sectPr，
    所以 restart=None 时必须**显式删除**继承来的 w:pgNumType，
    否则会把上一节的"页码重新开始"带过来。
    """
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sect_child(sec._sectPr, 'cols', num=str(cols), space=str(int(space_cm * 567)))
    sect_child(sec._sectPr, 'vAlign', val='center' if vcenter else 'top')
    if restart is not None:
        sect_child(sec._sectPr, 'pgNumType', start=str(restart))
    else:
        old = sec._sectPr.find(qn('w:pgNumType'))
        if old is not None:
            sec._sectPr.remove(old)
    if numbered:
        _numbered_footer(sec, footer_text)
    else:
        _blank_footer(sec)
    return sec


def init_cover_section(doc, cols=1, vcenter=True):
    """把文档的首节改造成封面节"""
    sec = doc.sections[0]
    sect_child(sec._sectPr, 'cols', num=str(cols), space=str(int(0.6 * 567)))
    sect_child(sec._sectPr, 'vAlign', val='center' if vcenter else 'top')
    _blank_footer(sec)
    return sec


# ────────────────────────── 三块内容 ──────────────────────────

def build_cover(doc, subject):
    """封面：学科 + 主标题 + 日期（垂直居中由节属性负责）"""
    _para(doc, '', Pt(12), after=Pt(24))
    _para(doc, subject, Pt(30), bold=True, after=Pt(10))
    _para(doc, BRAND_MAIN, Pt(40), bold=True, after=Pt(18))
    _para(doc, BRAND_DATE, Pt(14))


def build_colophon(doc, subject):
    """说明页（免责声明 / 凡例）"""
    _para(doc, COLOPHON_TITLE, Pt(16), bold=True, after=Pt(18))
    for line in COLOPHON_LINES:
        txt = line.format(subject=subject) if '{subject}' in line else line
        _para(doc, txt, Pt(11), align=WD_ALIGN_PARAGRAPH.LEFT, after=Pt(3), line=1.4)


def build_version_page(doc, subject, stats):
    """
    版本记录页
      stats = [(项目名, 规模说明), ...]  例如 [('大本解析', '75 章 · 1141 条目'), ...]
    """
    _para(doc, VERSION_TITLE, Pt(16), bold=True, after=Pt(20))
    _para(doc, 'v1.0　　2026 年 9 月 26 日　　首次整理', Pt(11),
          align=WD_ALIGN_PARAGRAPH.LEFT, after=Pt(14))
    _para(doc, '本册规模', Pt(11), bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=Pt(4))
    for name, desc in stats:
        _para(doc, '　%s　　%s' % (name, desc), Pt(11),
              align=WD_ALIGN_PARAGRAPH.LEFT, after=Pt(3))
    _para(doc, '', Pt(11), after=Pt(16))
    _para(doc, '整理排版：高三25班热心同学', Pt(11),
          align=WD_ALIGN_PARAGRAPH.LEFT, after=Pt(3))
    _para(doc, '内容版权归原出版方所有，本册不对内容主张任何权利。', Pt(11),
          align=WD_ALIGN_PARAGRAPH.LEFT, after=Pt(3))
    _para(doc, '问题反馈：%s' % GITHUB_URL, Pt(10),
          align=WD_ALIGN_PARAGRAPH.LEFT)
