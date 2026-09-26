#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
merge_zm.py — 周末必刷解析单独成册（小本）：
从原始 docx 提取 22 篇周末必刷的答案解析，排版同新整合（目录 + 居中章节标题 + 考点分段）。
输出: 新整合_周末必刷/周末必刷解析合集.docx（只读原始，不改任何源文件）
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

import merge_jiexi as MJ
import merge_book_order as MBO

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, '新整合_周末必刷')
OUT_FILE = os.path.join(OUT, '周末必刷解析合集.docx')

ZM_RE = re.compile(r'^周末必刷(\d+)')


def collect_zm():
    """收集 22 篇周末必刷（编号 -> 路径），并记住所属大概念"""
    items = []
    for name in sorted(os.listdir(ROOT)):
        d = os.path.join(ROOT, name)
        if not os.path.isdir(d) or not name.startswith('大概念'):
            continue
        for f in os.listdir(d):
            m = ZM_RE.match(f) if f.endswith('.docx') and not f.startswith('~$') and not f.endswith('_解析.docx') else None
            if m:
                items.append((int(m.group(1)), name, os.path.join(d, f)))
    items.sort(key=lambda x: x[0])
    return items


def main():
    os.makedirs(OUT, exist_ok=True)
    items = collect_zm()
    if not items:
        print('未找到周末必刷文档！')
        return

    doc = Document()
    MBO.build_page(doc)
    MBO.add_heading(doc, '周末必刷 解析合集', Pt(16), Pt(0), Pt(8))

    # 目录
    MBO.add_heading(doc, '目　录', Pt(16), Pt(0), Pt(8))
    for i, (num, day, path) in enumerate(items, 1):
        title = f'周末必刷{num}'
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = MJ.LINE_SPACE
        try:
            p.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
        except Exception:
            pass
        r1 = p.add_run(f'{i:02d}. {title}')
        MJ.set_font(r1, MJ.BODY_FONT, MJ.BODY_SIZE)
        r2 = p.add_run('\t')
        MJ.set_font(r2, MJ.BODY_FONT, MJ.BODY_SIZE)
        r3 = p.add_run('00')
        MJ.set_font(r3, MJ.BODY_FONT, MJ.BODY_SIZE)
    doc.add_page_break()

    # 正文：每篇标题 + 考点分段内容
    for num, day, path in items:
        MBO.render_chapter(doc, f'周末必刷{num}', path)
    doc.save(OUT_FILE)
    print(f'生成: {OUT_FILE}  共 {len(items)} 篇')


if __name__ == '__main__':
    main()
