#!/usr/bin/env python3
"""
merge_jiexi.py — 区域标记法版本

用开始/结束标记划定保留区域：
  考点速览区域：从 (数字)...(×/√) 开始，收集连续的提示行，到非考点速览内容结束
  典型例题区域：题号后的答案/解析行

丢弃无编号前缀的孤立提示行（来自教学内容的填空题答案）。
"""

import re
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree


def set_font(run, name='宋体', size=None, bold=False):
    """统一设置字体（含东亚字体回退）"""
    run.font.name = name
    run.font.bold = bold
    if size:
        run.font.size = size
    # 东亚字体回退
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = etree.SubElement(rPr, qn('w:rFonts'))
    rFonts.set(qn('w:eastAsia'), name)

# ── 正则 ──
PTN_MAIN_Q   = re.compile(r'^(\d+)\.\s*')               # "1."（年份全/半角、无年份均可）
PTN_EX_Q     = re.compile(r'^[【\[]\s*[例练变]\s*\d+\s*[】\]]\s*')  # "[例1]…" / "【例1】…" / "[练4]…"
PTN_TF_Q     = re.compile(r'^(\(\d+\))')                # "(1)…"
PTN_TF_ANS   = re.compile(r'[\(（][×√][\)）]')          # (×) (√)
PTN_BARE_LABEL = re.compile(r'^(答案|解析|提示)[：:]?$')  # 单独成段的 答案/解析/提示

PTN_ANSWER   = re.compile(r'^答案\s*[　 ]')
PTN_ANALYSIS = re.compile(r'^解析\s*[　 ]')
PTN_PROMPT   = re.compile(r'^提示\s*[　 ]')

# 多段续行判定（与 process_docs 一致）
from process_docs import is_cont_line

# ── 拆分判断题归一化 ——
# 原文档中部分判断题被拆成两段：题面段 "(4)油橄榄…不发生变化。" + 来源标记段 "(2024·甘肃卷，1C)(×)"
RE_TF_MARK_P = re.compile(r'[×√]')
RE_TF_START_P = re.compile(r'^[（(]\s*\d+\s*[）)]')
RE_SRC_LINE = re.compile(r'^[\(（]\d{4}[^（(]{0,30}[\)）]\s*[\(（][×√][\)）]\s*$')


def normalize_split_judges(texts):
    """把“题面段 + 来源标记段”的拆分判断题合并成一行，返回新列表。"""
    out = []
    i = 0
    n = len(texts)
    while i < n:
        t = texts[i]
        if (i + 1 < n and RE_TF_START_P.match(t) and not RE_TF_MARK_P.search(t)
                and RE_SRC_LINE.match(texts[i + 1])):
            out.append(t + texts[i + 1])
            i += 2
        else:
            out.append(t)
            i += 1
    return out

# ── 排版参数 ──
MARGIN_TOP    = Cm(0.7)
MARGIN_BOTTOM = Cm(0.7)
MARGIN_LEFT   = Cm(1.5)
MARGIN_RIGHT  = Cm(1.5)
BODY_FONT     = '宋体'
BODY_SIZE     = Pt(10.5)
TITLE_SIZE    = Pt(13)
LINE_SPACE    = 1.15     # 行距倍数
PARA_AFTER    = Pt(1)    # 段后极小间距
BLOCK_GAP     = Pt(3)    # 条目间间距（考点速览块之间、典型例题之间）
PAGE_NUM_SIZE = Pt(12)    # 页码字号（页脚外侧页码，突出可读）


def add_page_number(paragraph, alignment):
    """在段落中添加 PAGE 域（自动页码），并设置对齐方式"""
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after  = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run()
    set_font(run, BODY_FONT, PAGE_NUM_SIZE)
    # w:fldChar type="begin"
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    run._element.append(fld_begin)
    # w:instrText PAGE
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    run._element.append(instr)
    # w:fldChar type="end"
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._element.append(fld_end)


def setup_page_numbers(doc):
    """设置交替页码：奇数页右下，偶数页左下"""
    section = doc.sections[0]
    sectPr = section._sectPr
    sectPr.set(qn('w:evenAndOddHeaders'), '1')
    # 奇数页页脚（右下）
    footer_odd = section.footer
    footer_odd.is_linked_to_previous = False
    for p in footer_odd.paragraphs:
        p.clear()
    if footer_odd.paragraphs:
        add_page_number(footer_odd.paragraphs[0], WD_ALIGN_PARAGRAPH.RIGHT)
    else:
        p = footer_odd.add_paragraph()
        add_page_number(p, WD_ALIGN_PARAGRAPH.RIGHT)
    # 偶数页页脚（左下）
    footer_even = section.even_page_footer
    footer_even.is_linked_to_previous = False
    for p in footer_even.paragraphs:
        p.clear()
    if footer_even.paragraphs:
        add_page_number(footer_even.paragraphs[0], WD_ALIGN_PARAGRAPH.LEFT)
    else:
        p = footer_even.add_paragraph()
        add_page_number(p, WD_ALIGN_PARAGRAPH.LEFT)

# ── 分类规则 ──
CATEGORY_RULES = [
    ('第X讲',     re.compile(r'^第\d+讲')),
    ('高考共鸣点', re.compile(r'^高考共鸣点')),
    ('阶段排查',   re.compile(r'^阶段排查')),
    ('周末必刷',   re.compile(r'^周末必刷')),
]


def classify_filename(fname):
    for cat, pat in CATEGORY_RULES:
        if pat.match(fname):
            return cat
    return None


# 题目来源（年份+卷别+题号）：(2025·浙江6月卷，18) / （2024·浙江卷，11)
RE_SRC = re.compile(r'[（(]\s*\d{4}[^（(）)]{0,25}[）)]')


# 题号/题干特征（用于区分“知识小标题 1.萨顿的假说”与“真题 1.(2024·…)”）
RE_Q_YEAR = re.compile(r'[（(]\s*\d{4}')
RE_Q_EMPTY_PAREN = re.compile(r'[（(]\s*[　 ]*[）)]')
RE_Q_WORDS = re.compile(r'下列|正确的是|错误的是|叙述|推断|符合|分析|计算|判断|回答|说明|探究|验证|假设|设计|图解|理由|原因')


def looks_like_question(text):
    """
    判断是否为题目行：
    - 无编号题：以年份开头，如 “(2022·江苏卷，11)摩尔根…”
    - 编号题：N. / N、/ N． 后带 年份/题干关键词/空括号（　　）
    知识小标题如 “1.萨顿的假说” 不算题。

    注意：题号后的标点【必须存在】。曾经的 `[.、．]?`（可选）会把题干预设被拆行后
    以「数字+空格」开头的续段误判成新题号（如 “3 500 m的高山竹林中…(　　)”
    被判成题号 3.、“240 μmol·L－1镉离子浓度下…” 被判成题号 240.），
    导致前一道真题的题号/来源被截断。
    """
    if re.match(r'^[（(]\s*\d{4}', text):
        return True
    m = re.match(r'^(\d+)[.、．]\s*(.*)$', text)
    if not m:
        return False
    rest = m.group(2)
    if RE_Q_YEAR.search(text):
        return True
    if RE_Q_EMPTY_PAREN.search(text):
        return True
    if RE_Q_WORDS.search(rest) and len(rest) >= 12:
        return True
    return False


def extract_qnum(text):
    """
    提取题号：'1.' / ''（无编号题） / '【例1】' / '(1)'。
    知识小标题（1.萨顿的假说）、非题目行返回 None。
    """
    if looks_like_question(text):
        if re.match(r'^[（(]\s*\d{4}', text):
            return ''  # 无编号题（如“考向”下 (2022·江苏卷，11)…）
        m = re.match(r'^(\d+)[.、．]\s*', text)
        if m:
            return m.group(1) + '.'
    m = PTN_EX_Q.match(text)
    if m:
        return m.group(0).strip()
    m = PTN_TF_Q.match(text)
    if m and PTN_TF_ANS.search(text):
        return m.group(1)
    return None


def is_keeper(text):
    return bool(PTN_ANSWER.match(text) or PTN_ANALYSIS.match(text))


def process_one_docx(input_path):
    """
    用区域标记法提取条目。

    返回: [{'num': '1.' 或 '', 'lines': [行1, 行2, ...]}]
    """
    doc = Document(input_path)
    # 提取所有非空段落文本（先合并被拆成两段的判断题）
    texts = normalize_split_judges([p.text.strip() for p in doc.paragraphs if p.text.strip()])

    entries = []
    i = 0
    n = len(texts)

    while i < n:
        t = texts[i]

        # ── 考点速览区域 ──
        # 以 (数字)...(×/√) 开头 → 保留完整行 + 后续提示/答案/解析（含其续段）
        if PTN_TF_Q.match(t) and PTN_TF_ANS.search(t):
            block = [t]
            i += 1
            while i < n:
                nt = texts[i]
                if PTN_PROMPT.match(nt) or PTN_ANSWER.match(nt) or PTN_ANALYSIS.match(nt):
                    block.append(nt)
                    i += 1
                    # 多段答案/解析/提示：后续 (2)/②… 或句子续文
                    while i < n and is_cont_line(block[-1], texts[i]):
                        block.append(texts[i])
                        i += 1
                else:
                    break
            entries.append({'num': '', 'lines': block})
            continue

        # ── 主考题区域（含无年份题、全角年份题、无编号题）──
        qnum = extract_qnum(t)
        if qnum is not None:
            m_src = RE_SRC.search(t)
            src = m_src.group(0) if m_src else ''
            i += 1
            keeper_lines = []
            while i < n:
                nt = texts[i]
                if PTN_ANSWER.match(nt) or PTN_ANALYSIS.match(nt):
                    keeper_lines.append(nt)
                    i += 1
                    # 多段答案/解析：后续 (2)/②… 或句子续文
                    while i < n and is_cont_line(keeper_lines[-1], texts[i]):
                        keeper_lines.append(texts[i])
                        i += 1
                elif PTN_BARE_LABEL.match(nt) and keeper_lines:
                    # 单独成段的 答案/解析/提示：保留标签并收集其后续内容
                    keeper_lines.append(nt)
                    i += 1
                    while i < n and is_cont_line(keeper_lines[-1], texts[i]):
                        keeper_lines.append(texts[i])
                        i += 1
                elif PTN_PROMPT.match(nt):
                    # 提示行不跟主题目，跳过（防止知识点提示混入）
                    i += 1
                elif extract_qnum(nt) is not None or (PTN_TF_Q.match(nt) and PTN_TF_ANS.search(nt)):
                    break  # 新题号开始
                else:
                    i += 1  # 跳过题目正文、选项等
            if keeper_lines:
                entries.append({'num': qnum, 'src': src, 'lines': keeper_lines})
            continue

        # 孤立答案/解析兜底：没有题目编号跟随的答案/解析（如连线题答案）保留为无编号条目
        if PTN_ANSWER.match(t) or PTN_ANALYSIS.match(t):
            lines = [t]
            i += 1
            while i < n and is_cont_line(lines[-1], texts[i]):
                lines.append(texts[i])
                i += 1
            entries.append({'num': '', 'src': '', 'lines': lines})
            continue

        # 其他行（知识点、孤立提示等）→ 丢弃
        i += 1

    return entries


def merge_category(daginian_dir, daginian_name, cat_name, filepaths):
    """合并一个分类下所有文件的条目，紧凑排版输出到 解析合集/"""
    out_dir = os.path.join('解析合集', daginian_name)
    os.makedirs(out_dir, exist_ok=True)
    
    merged = Document()
    
    # ── 页面边距 ──
    for section in merged.sections:
        section.top_margin = MARGIN_TOP
        section.bottom_margin = MARGIN_BOTTOM
        section.left_margin = MARGIN_LEFT
        section.right_margin = MARGIN_RIGHT
        section.header_distance = Cm(0)
        section.footer_distance = Cm(0.3)
    
    # ── 全局默认样式 ──
    style = merged.styles['Normal']
    style.font.name = BODY_FONT
    style.font.size = BODY_SIZE
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = PARA_AFTER
    style.paragraph_format.line_spacing = LINE_SPACE
    
    # 交替页码
    setup_page_numbers(merged)
    
    # 大标题
    tp = merged.add_paragraph()
    tp.paragraph_format.space_before = Pt(4)
    tp.paragraph_format.space_after = Pt(6)
    tp.paragraph_format.line_spacing = 1.0
    run = tp.add_run(f'{daginian_name} —— {cat_name}')
    set_font(run, BODY_FONT, Pt(14), bold=True)
    
    for fpath in sorted(filepaths):
        fname = os.path.basename(fpath)
        name_no_ext = fname.replace('.docx', '')
        
        # 小节标题
        st = merged.add_paragraph()
        st.paragraph_format.space_before = Pt(8)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.0
        run = st.add_run(name_no_ext)
        set_font(run, BODY_FONT, TITLE_SIZE, bold=True)
        
        entries = process_one_docx(fpath)
        if not entries:
            np = merged.add_paragraph('（无条目）')
            np.paragraph_format.space_before = Pt(0)
            np.paragraph_format.space_after = Pt(0)
            continue
        
        first_entry = True
        for entry in entries:
            lines = entry['lines']
            if not lines:
                continue
            
            # 条目间分隔间距（除第一个外）
            if not first_entry:
                gap = merged.add_paragraph()
                gap.paragraph_format.space_before = Pt(0)
                gap.paragraph_format.space_after = Pt(0)
                gap.paragraph_format.line_spacing = BLOCK_GAP / Pt(12)
                gr = gap.add_run('')
                gr.font.size = Pt(1)
            first_entry = False
            
            if entry['num'] == '' and not entry.get('src'):
                # 考点速览块：每行独立段落
                for line in lines:
                    ep = merged.add_paragraph(line)
                    ep.paragraph_format.space_before = Pt(0)
                    ep.paragraph_format.space_after = Pt(0)
                    ep.paragraph_format.line_spacing = LINE_SPACE
                    for run in ep.runs:
                        set_font(run, BODY_FONT, BODY_SIZE)
            else:
                # 典型例题：题号+来源 + 首行内容同行
                first = lines[0]
                p = merged.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = LINE_SPACE
                head = (entry['num'] or '') + (entry.get('src') or '')
                if head:
                    rn = p.add_run(head + ' ')
                else:
                    rn = p.add_run('')
                set_font(rn, BODY_FONT, BODY_SIZE, bold=True)
                r1 = p.add_run(first)
                set_font(r1, BODY_FONT, BODY_SIZE)
                
                # 后续行（通常是解析续行）
                for extra in lines[1:]:
                    ep = merged.add_paragraph(extra)
                    ep.paragraph_format.space_before = Pt(0)
                    ep.paragraph_format.space_after = Pt(0)
                    ep.paragraph_format.line_spacing = LINE_SPACE
                    for run in ep.runs:
                        set_font(run, BODY_FONT, BODY_SIZE)
    
    out_path = os.path.join(out_dir, f'{cat_name}.docx')
    try:
        merged.save(out_path)
    except PermissionError:
        # 目标文件被占用（如Word打开中），改用带后缀的文件名
        out_path = os.path.join(out_dir, f'{cat_name}_页码.docx')
        merged.save(out_path)
    print(f'  {cat_name}: {out_path}')


def process_daginian(root_dir):
    daginian_name = os.path.basename(root_dir)
    
    categorized = {cat: [] for cat in ['第X讲', '高考共鸣点', '阶段排查', '周末必刷']}
    for fname in os.listdir(root_dir):
        if not fname.endswith('.docx') or fname.startswith('~$') or fname.endswith('_解析.docx'):
            continue
        cat = classify_filename(fname)
        if cat:
            categorized[cat].append(os.path.join(root_dir, fname))
    
    for cat, files in categorized.items():
        if files:
            merge_category(root_dir, daginian_name, cat, files)


def main():
    roots = sorted(d for d in os.listdir('.') if os.path.isdir(d) and d.startswith('大概念'))
    if not roots:
        print('未找到大概念文件夹！')
        return
    
    os.makedirs('解析合集', exist_ok=True)
    for root in roots:
        print(f'\n处理: {root}')
        process_daginian(root)
    print('\n全部完成！输出在 解析合集/ 文件夹')


if __name__ == '__main__':
    main()
