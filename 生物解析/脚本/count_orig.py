# -*- coding: utf-8 -*-
"""tailwind_orig.py — 原始文档内容构成：正文/表格/图片（只读统计）"""
import os
from docx import Document

ROOT = os.path.dirname(os.path.abspath(__file__))

para_chars = 0
table_chars = 0
n_tables = 0
n_images = 0          # 文档内嵌图片（inline shapes）数
n_tf_ref = 0           # 段落中的图片引用行 [F:...]
n_docs = 0
files_list = []

for name in sorted(os.listdir(ROOT)):
    d = os.path.join(ROOT, name)
    if not os.path.isdir(d) or not name.startswith('大概念'):
        continue
    for f in sorted(os.listdir(d)):
        if not f.endswith('.docx') or f.startswith('~$') or f.endswith('_解析.docx'):
            continue
        files_list.append(f)
        n_docs += 1
        doc = Document(os.path.join(d, f))
        for p in doc.paragraphs:
            para_chars += len(p.text)
            if p.text.strip().startswith('[F:'):
                n_tf_ref += 1
        for tbl in doc.tables:
            n_tables += 1
            seen = set()
            for row in tbl.rows:
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    for p in cell.paragraphs:
                        table_chars += len(p.text)
        try:
            n_images += len(doc.inline_shapes)
        except Exception:
            pass

lines = []
lines.append(f'原始文档数: {n_docs}')
lines.append(f'正文段落字符: {para_chars:,}')
lines.append(f'表格数量: {n_tables:,}  表格内字符: {table_chars:,}')
lines.append(f'正文图片引用行 [F:...]: {n_tf_ref:,}')
lines.append(f'文档内嵌图片/图形对象: {n_images:,}')
lines.append('')
lines.append('注意：python-docx 只能数“正文段落”和“顶层表格”；')
lines.append('    大量插图以图片引用行 [F:xxx] 或内嵌对象存在，无法统计其视觉面积。')

with open('orig_composition.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines) + '\n')
print('\n'.join(lines))
