# -*- coding: utf-8 -*-
"""probe21.py — 第21讲题号/截断问题定位（只读）"""
import re
from docx import Document

P = '大概念四　遗传信息控制生物性状并代代相传——遗传规律与人类遗传病\\第21讲　基因在染色体上、伴性遗传.docx'
BOOK = '新整合_全书\\解析合集_全书.docx'


def dump_orig():
    ps = [x.text.strip() for x in Document(P).paragraphs if x.text.strip()]
    out = [f'原文段落数: {len(ps)}', '=' * 70]
    for i, t in enumerate(ps):
        tag = ''
        if re.match(r'^考点', t):
            tag = '<考点>'
        elif re.match(r'^考向', t):
            tag = '<考向>'
        elif re.match(r'^\d+\.\s*[（(]\d{4}', t):
            tag = '<题>'
        elif re.match(r'^答案', t):
            tag = '<答案>'
        elif re.match(r'^解析', t):
            tag = '<解析>'
        elif re.match(r'^[（(]\d+[）)]', t) and re.search(r'[×√]', t):
            tag = '<判断>'
        out.append(f'#{i:03d} {tag} {t[:95]}')
    return '\n'.join(out)


def dump_book():
    doc = Document(BOOK)
    ps = [x.text.strip() for x in doc.paragraphs if x.text.strip()]
    start = None
    for i, t in enumerate(ps):
        if t == '第21讲　基因在染色体上、伴性遗传':
            start = i
            break
    out = [f'全书第21讲起始段: {start}', '=' * 70]
    for i, t in enumerate(ps[start:start + 70]):
        out.append(f'+{i:03d} {t[:95]}')
    return '\n'.join(out)


with open('probe21.txt', 'w', encoding='utf-8') as f:
    f.write(dump_orig() + '\n\n' + dump_book() + '\n')
print('done')
